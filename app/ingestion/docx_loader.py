"""DOCX document loader and text extractor using python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import List
import docx
from app.utils.logging import get_logger

logger = get_logger("app.ingestion.docx_loader")


class DOCXLoader:
    """Extracts paragraphs, headings, and tables from Word (.docx) documents."""

    def __init__(self, file_path: str | Path) -> None:
        self.path = Path(file_path)

    def load(self) -> List[dict]:
        """Extract content from the target DOCX file.

        Returns:
            List of dicts representing structured pages/sections:
                - document: filename
                - page: approximate page number (1-indexed)
                - text: extracted text
                - section: detected section name
        """
        if not self.path.exists():
            logger.error(f"DOCX file not found: {self.path}")
            raise FileNotFoundError(f"DOCX file not found: {self.path}")

        try:
            doc = docx.Document(str(self.path))
        except Exception as exc:
            logger.error(f"Failed to open or parse DOCX '{self.path.name}': {exc}")
            raise ValueError(f"Corrupted or invalid DOCX file '{self.path.name}': {exc}") from exc

        extracted_pages: List[dict] = []
        current_section = "General"
        current_text_buffer: List[str] = []
        current_page = 1
        word_count_in_page = 0

        # Process paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = getattr(para.style, "name", "").lower() if para.style else ""
            is_heading = style_name.startswith("heading") or style_name in ["title", "subtitle"]

            if is_heading:
                # If we had accumulated substantial content, flush it as a page/section block
                if current_text_buffer and word_count_in_page > 150:
                    extracted_pages.append({
                        "document": self.path.name,
                        "document_path": str(self.path),
                        "page": current_page,
                        "text": "\n".join(current_text_buffer),
                        "section": current_section,
                    })
                    current_page += 1
                    current_text_buffer = []
                    word_count_in_page = 0

                current_section = text
                current_text_buffer.append(f"## {text}")
            else:
                current_text_buffer.append(text)
                word_count_in_page += len(text.split())

                # Logical page break every ~450 words if no heading intervened
                if word_count_in_page >= 450:
                    extracted_pages.append({
                        "document": self.path.name,
                        "document_path": str(self.path),
                        "page": current_page,
                        "text": "\n".join(current_text_buffer),
                        "section": current_section,
                    })
                    current_page += 1
                    current_text_buffer = []
                    word_count_in_page = 0

        # Process tables
        for table in doc.tables:
            table_lines: List[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                # Filter out adjacent duplicate cells (common in merged cells)
                unique_cells = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                if any(unique_cells):
                    table_lines.append(" | ".join(unique_cells))

            if table_lines:
                table_text = "\n".join(table_lines)
                current_text_buffer.append(f"\n[Table Data]:\n{table_text}\n")
                word_count_in_page += len(table_text.split())

        # Flush any remaining buffer
        if current_text_buffer:
            extracted_pages.append({
                "document": self.path.name,
                "document_path": str(self.path),
                "page": current_page,
                "text": "\n".join(current_text_buffer),
                "section": current_section,
            })

        logger.info(f"Extracted {len(extracted_pages)} logical page(s) from DOCX: {self.path.name}")
        return extracted_pages

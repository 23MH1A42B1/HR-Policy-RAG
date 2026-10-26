"""PPTX presentation loader and slide text extractor using python-pptx."""

from __future__ import annotations

from pathlib import Path
from typing import List
from pptx import Presentation
from app.utils.logging import get_logger

logger = get_logger("app.ingestion.pptx_loader")


class PPTXLoader:
    """Extracts text slide-by-slide from PowerPoint (.pptx) presentations."""

    def __init__(self, file_path: str | Path) -> None:
        self.path = Path(file_path)

    def load(self) -> List[dict]:
        """Extract slides and their content from the target PPTX file.

        Returns:
            List of dicts representing slides:
                - document: filename
                - page: 1-indexed slide number
                - text: text content extracted from slide shapes and tables
                - section: detected slide title
        """
        if not self.path.exists():
            logger.error(f"PPTX file not found: {self.path}")
            raise FileNotFoundError(f"PPTX file not found: {self.path}")

        try:
            prs = Presentation(str(self.path))
        except Exception as exc:
            logger.error(f"Failed to open or parse PPTX '{self.path.name}': {exc}")
            raise ValueError(f"Corrupted or invalid PPTX file '{self.path.name}': {exc}") from exc

        extracted_slides: List[dict] = []
        total_slides = len(prs.slides)
        logger.info(f"Extracting {total_slides} slide(s) from PPTX: {self.path.name}")

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = ""
            slide_text_elements: List[str] = []

            # 1. Identify slide title if present
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text.strip()
                if slide_title:
                    slide_text_elements.append(f"## {slide_title}")

            # 2. Iterate through all shapes on slide
            for shape in slide.shapes:
                # Skip the title shape as it was already recorded
                if shape == slide.shapes.title:
                    continue

                # Text frames
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_text_elements.append(text)

                # Tables
                elif shape.has_table:
                    table = shape.table
                    table_rows = []
                    for row in table.rows:
                        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text.strip()]
                        if row_text:
                            table_rows.append(" | ".join(row_text))
                    if table_rows:
                        slide_text_elements.append("[Table]:\n" + "\n".join(table_rows))

            # 3. Slide notes if present
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_text_elements.append(f"[Presenter Notes]: {notes_text}")

            # Fallback title from first text element if title shape was not explicitly configured
            if not slide_title and slide_text_elements:
                first_elem = slide_text_elements[0].replace("## ", "")
                if len(first_elem) <= 80 and not first_elem.startswith(("[", "•", "-", "*")):
                    slide_title = first_elem

            combined_text = "\n\n".join(slide_text_elements).strip()
            section_name = slide_title if slide_title else f"Slide {slide_idx}"

            extracted_slides.append({
                "document": self.path.name,
                "document_path": str(self.path),
                "page": slide_idx,
                "text": combined_text,
                "section": section_name,
            })

        return extracted_slides

"""Generates realistic sample HR policy documents (PDF, DOCX, PPTX).

Fictional company: "Acme Global Enterprise"
Effective year: 2026
"""

from __future__ import annotations

from pathlib import Path
import docx
from docx.shared import RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches as PptxInches, Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


def generate_pdf_leave_policy(output_path: Path) -> None:
    """Generate sample leave_policy.pdf using ReportLab."""
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#312E81"),
        spaceAfter=14,
    )
    h1_style = ParagraphStyle(
        "DocH1",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#4338CA"),
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=8,
    )

    story = [
        Paragraph("Acme Global Enterprise", body_style),
        Paragraph("Comprehensive Employee Leave Policy (Rev. 2026)", title_style),
        Paragraph(
            "<b>Policy ID:</b> HR-POL-2026-01 | <b>Effective Date:</b> January 1, 2026 | <b>Scope:</b> All Full-Time Employees",
            body_style,
        ),
        Spacer(1, 10),
        Paragraph("1. Annual Leave Entitlement", h1_style),
        Paragraph(
            "All regular full-time employees are entitled to 20 business days of paid annual leave per calendar year. "
            "Leave accrues monthly at a rate of 1.67 days per completed month of service starting from the first day of employment. "
            "Employees are encouraged to take their annual leave to maintain work-life balance and overall wellness.",
            body_style,
        ),
        Paragraph(
            "Leave requests exceeding 5 consecutive business days must be submitted through the HR Portal at least 14 days in advance. "
            "Managerial approval is required based on operational coverage.",
            body_style,
        ),
        Paragraph("2. Carryover Rules and Expiration", h1_style),
        Paragraph(
            "A maximum of 5 unused annual leave days may be carried over into the following calendar year. "
            "All carried-over days must be utilized before March 31st of the new calendar year, after which unused carryover days will forfeit without cash compensation.",
            body_style,
        ),
        PageBreak(),
        Paragraph("3. Sick Leave & Medical Absence", h1_style),
        Paragraph(
            "Employees receive 10 paid sick days annually, available at the beginning of each calendar year. "
            "Sick leave covers personal illness, routine medical examinations, and caring for an immediate family member who requires assistance. "
            "If an employee is absent due to illness for 3 or more consecutive working days, a signed certificate from a licensed healthcare provider is mandatory upon returning.",
            body_style,
        ),
        Paragraph("4. Parental and Family Leave", h1_style),
        Paragraph(
            "Acme Global provides 16 consecutive weeks of fully paid parental leave for primary caregivers following the birth, adoption, or foster placement of a child. "
            "Secondary caregivers are entitled to 6 consecutive weeks of fully paid leave. "
            "Parental leave must be taken within 12 months following the qualifying life event.",
            body_style,
        ),
        PageBreak(),
        Paragraph("5. Bereavement and Compassionate Leave", h1_style),
        Paragraph(
            "Employees are eligible for up to 5 consecutive paid business days in the unfortunate event of the death of an immediate family member "
            "(spouse, child, parent, sibling, or domestic partner). For extended relatives (grandparents, in-laws), up to 3 paid business days are granted.",
            body_style,
        ),
        Paragraph("6. Leave Request Submission Workflow", h1_style),
        Paragraph(
            "All time-off applications must be submitted digitally via the Acme HR Employee Self-Service Portal. "
            "Your direct supervisor has 48 business hours to review and take action on the request. "
            "For urgent or unplanned absences, please notify your team lead via email or Slack prior to 9:00 AM local time.",
            body_style,
        ),
    ]

    doc.build(story)


def generate_docx_remote_work_policy(output_path: Path) -> None:
    """Generate sample remote_work_policy.docx using python-docx."""
    doc = docx.Document()

    # Title
    title = doc.add_heading("Remote and Hybrid Work Policy 2026", level=0)
    title.runs[0].font.color.rgb = RGBColor(49, 46, 129)

    doc.add_paragraph("Company: Acme Global Enterprise | Document Version: 2.4 (2026)")
    doc.add_paragraph(
        "This policy defines the standards, eligibility criteria, and expectations for employees working in remote and hybrid arrangements."
    )

    # Section 1
    doc.add_heading("1. Eligibility and Work Models", level=1)
    doc.add_paragraph(
        "Employees whose roles do not require physical presence on company premises may be eligible for hybrid or full remote arrangements, "
        "subject to department head approval and sustained satisfactory performance evaluations."
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Model"
    hdr_cells[1].text = "Onsite Days"
    hdr_cells[2].text = "Eligible Roles"

    row1 = table.add_row().cells
    row1[0].text = "Hybrid (Standard)"
    row1[1].text = "2 - 3 Days / Week"
    row1[2].text = "General Engineering, Marketing, Finance"

    row2 = table.add_row().cells
    row2[0].text = "Fully Remote"
    row2[1].text = "0 Days (Quarterly offsites)"
    row2[2].text = "Designated Distributed Teams"

    row3 = table.add_row().cells
    row3[0].text = "Onsite Priority"
    row3[1].text = "5 Days / Week"
    row3[2].text = "Facilities, Hardware Labs, Office Operations"

    # Section 2
    doc.add_heading("2. Core Working Hours and Availability", level=1)
    doc.add_paragraph(
        "To ensure seamless cross-functional collaboration, all remote and hybrid employees must be available and reachable during core collaboration hours "
        "between 10:00 AM and 4:00 PM in their designated primary time zone. "
        "Team members must maintain an updated Slack status and responsive calendar."
    )

    # Section 3
    doc.add_heading("3. Home Office Stipend and Ergonomic Equipment", level=1)
    doc.add_paragraph(
        "Eligible full-time remote and hybrid employees receive a one-time reimbursement of up to $500 for ergonomic office furniture, "
        "monitors, keyboards, or desk setups. Additionally, employees receive a recurring monthly allowance of $50 toward high-speed home internet service. "
        "All expense claims must be filed through Expensify within 45 days of purchase."
    )

    # Section 4
    doc.add_heading("4. Cybersecurity and Information Protection", level=1)
    doc.add_paragraph(
        "Remote staff must strictly adhere to Acme Global's Security Baseline. Company-issued laptops must use full-disk encryption, "
        "passwords must comply with MFA requirements, and connecting to company resources from unsecured public Wi-Fi without the Corporate VPN is strictly forbidden. "
        "Confidential documents must never be printed on unauthorized home printers."
    )

    doc.save(str(output_path))


def generate_pptx_expense_policy(output_path: Path) -> None:
    """Generate sample expense_policy.pptx using python-pptx."""
    prs = Presentation()
    prs.slide_width = PptxInches(10)
    prs.slide_height = PptxInches(5.625)

    blank_layout = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    # Slide 1: Title
    slide1 = prs.slides.add_slide(title_layout)
    slide1.shapes.title.text = "Corporate Travel & Expense Policy 2026"
    subtitle = slide1.placeholders[1]
    subtitle.text = "Acme Global Enterprise | Finance & People Operations (Rev. 2026)"

    # Helper function to add a standard slide
    def add_content_slide(title_text: str, bullets: list[str], notes: str = ""):
        slide = prs.slides.add_slide(blank_layout)
        # Title box
        tx_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.5), PptxInches(8.4), PptxInches(0.8))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = PptxPt(22)
        p.font.bold = True
        p.font.color.rgb = PptxRGBColor(49, 46, 129)

        # Content box
        content_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(1.4), PptxInches(8.4), PptxInches(3.5))
        c_tf = content_box.text_frame
        c_tf.word_wrap = True
        for idx, bullet in enumerate(bullets):
            p_bullet = c_tf.paragraphs[0] if idx == 0 else c_tf.add_paragraph()
            p_bullet.text = f"• {bullet}"
            p_bullet.font.size = PptxPt(14)
            p_bullet.space_after = PptxPt(10)

        if notes:
            notes_slide = slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = notes
        return slide

    # Slide 2: Travel & Lodging
    add_content_slide(
        "1. Business Travel and Hotel Accommodations",
        [
            "All business flights under 6 hours must be booked in Economy or Premium Economy class.",
            "Domestic hotel accommodations must not exceed $220 per night (exclusive of local taxes), unless traveling to high-cost metropolitan tier-1 cities where the cap is $320.",
            "All flight and rail reservations must be processed through the Navan corporate portal at least 14 days prior to departure.",
        ],
        notes="Remind employees that luxury upgrades are at personal expense.",
    )

    # Slide 3: Meal Allowances
    add_content_slide(
        "2. Daily Meal Allowances (Per Diem)",
        [
            "Breakfast allowance: up to $20 per day (when traveling overnight).",
            "Lunch allowance: up to $30 per day for traveling team members.",
            "Dinner allowance: up to $50 per day for overnight business travel.",
            "Itemized receipts are required for all meals above $25. Alcoholic beverages are non-reimbursable unless attending an official executive client dinner.",
        ],
    )

    # Slide 4: Non-Reimbursable Expenses
    add_content_slide(
        "3. Non-Reimbursable Personal Expenses",
        [
            "Traffic and parking violations, speeding tickets, or toll violations incurred while driving.",
            "In-flight personal entertainment, Wi-Fi upgrades on short flights (< 2 hours), or streaming subscriptions.",
            "Personal grooming, clothing, hotel minibar charges, dry cleaning for trips under 4 days.",
            "Spouse or guest travel expenses unless specifically authorized by the VP of People.",
        ],
    )

    # Slide 5: Expense Reporting & Deadlines
    add_content_slide(
        "4. Submission Deadlines and Approval Limits",
        [
            "Expense reports must be submitted within 30 calendar days of the expenditure date.",
            "Expenses below $1,000 require Direct Manager sign-off.",
            "Expenses between $1,000 and $5,000 require Department Director approval.",
            "Expenses exceeding $5,000 require VP or CFO approval prior to disbursement.",
        ],
        notes="Reimbursements are processed bi-weekly with payroll.",
    )

    prs.save(str(output_path))


def generate_pdf_attendance_policy(output_path: Path) -> None:
    """Generate sample attendance_policy.pdf using ReportLab."""
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "AttTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1E1B4B"),
        spaceAfter=14,
    )
    h1_style = ParagraphStyle(
        "AttH1",
        parent=styles["Heading2"],
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#4338CA"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "AttBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=8,
    )

    story = [
        Paragraph("Acme Global Enterprise", body_style),
        Paragraph("Working Hours, Attendance, and Overtime Policy (2026)", title_style),
        Paragraph("<b>Policy ID:</b> HR-ATT-2026 | <b>Effective:</b> 2026 | <b>Applicability:</b> Global Workforce", body_style),
        Spacer(1, 10),
        Paragraph("1. Standard Working Hours", h1_style),
        Paragraph(
            "The standard full-time workweek comprises 40 hours, organized as five 8-hour days, Monday through Friday, 9:00 AM to 5:00 PM local time. "
            "Flexible work schedules may be established with prior agreement between the employee and their department manager, provided business coverage is preserved.",
            body_style,
        ),
        Paragraph("2. Punctuality and Unplanned Absences", h1_style),
        Paragraph(
            "Employees are expected to arrive punctually and be ready to begin work at their scheduled start time. "
            "In cases of unexpected emergency or illness, notice must be communicated to the team lead via phone or messaging at least 30 minutes before the scheduled start time.",
            body_style,
        ),
        PageBreak(),
        Paragraph("3. Overtime Regulations and Authorization", h1_style),
        Paragraph(
            "Non-exempt employees who work in excess of 40 hours in a given workweek will receive overtime compensation at 1.5 times their regular hourly rate. "
            "All overtime must be pre-authorized in writing by the direct supervisor. Working overtime without prior authorization is strictly prohibited.",
            body_style,
        ),
        Paragraph("4. Attendance Records and Progressive Discipline", h1_style),
        Paragraph(
            "Accurate timecards must be submitted by Friday 5:00 PM every week. "
            "Patterns of unexcused tardiness, unauthorized absence, or failure to record working hours will lead to progressive disciplinary steps: "
            "(1) Informal Coaching, (2) Written Warning, (3) Performance Improvement Plan, and (4) Termination of Employment.",
            body_style,
        ),
    ]

    doc.build(story)


def main() -> None:
    raw_dir = Path("data/raw")
    raw_dir.mkdir(parents=True, exist_ok=True)

    print("Generating sample policy documents for development...")

    p1 = raw_dir / "leave_policy.pdf"
    generate_pdf_leave_policy(p1)
    print(f"Created: {p1}")

    p2 = raw_dir / "remote_work_policy.docx"
    generate_docx_remote_work_policy(p2)
    print(f"Created: {p2}")

    p3 = raw_dir / "expense_policy.pptx"
    generate_pptx_expense_policy(p3)
    print(f"Created: {p3}")

    p4 = raw_dir / "attendance_policy.pdf"
    generate_pdf_attendance_policy(p4)
    print(f"Created: {p4}")

    print("All sample policy documents generated successfully.")


if __name__ == "__main__":
    main()
